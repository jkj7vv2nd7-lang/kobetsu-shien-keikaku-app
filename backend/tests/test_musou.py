"""無双化テスト群 v0.9（追記専用。実行: python -m pytest backend/tests -q）"""
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE / "backend"))
sys.path.insert(0, str(BASE / "src"))


def _client():
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    auth_mod.ensure_user("m_teacher", "m_teacher123", "teacher", "")
    auth_mod.ensure_user("m_manager", "m_manager123", "manager", "")
    # 他テストの影響を受けないよう確定パスワードに揃える
    with auth_mod.db.conn() as _c:
        _t = _c.execute("SELECT id FROM users WHERE username='m_teacher'").fetchone()
        _m = _c.execute("SELECT id FROM users WHERE username='m_manager'").fetchone()
    auth_mod.set_password(dict(_t)["id"], "m_teacher123")
    auth_mod.set_password(dict(_m)["id"], "m_manager123")
    cc = TestClient(app)
    t = cc.post("/api/auth/login", json={"username": "m_teacher", "password": "m_teacher123"}).json()
    m = cc.post("/api/auth/login", json={"username": "m_manager", "password": "m_manager123"}).json()
    return cc, {"Authorization": f"Bearer {t['token']}"}, {"Authorization": f"Bearer {m['token']}"}


def test_propose_plain_no_double_suffix():
    sys.path.insert(0, str(BASE / "src"))
    from core.plain import propose_plain
    r = propose_plain("通級指導教室と連携し、通級でも様子を見る。")
    assert "指導教室指導教室" not in r["rewritten"]
    assert "通級指導教室と連携" in r["rewritten"]


def test_must_change_pw_flow():
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    auth_mod.ensure_user("m_fresh", "m_fresh123", "teacher", "")
    with auth_mod.db.conn() as _c:
        _r = _c.execute("SELECT id FROM users WHERE username='m_fresh'").fetchone()
    auth_mod.set_password(dict(_r)["id"], "m_fresh123")
    with auth_mod.db.conn() as _c2:
        _c2.execute("UPDATE users SET must_change_pw=1 WHERE username='m_fresh'")
    cc = TestClient(app)
    t = cc.post("/api/auth/login", json={"username": "m_fresh", "password": "m_fresh123"}).json()
    assert t.get("must_change_pw") is True
    th = {"Authorization": f"Bearer {t['token']}"}
    me = cc.get("/api/auth/me", headers=th).json()
    assert me["must_change_pw"] is True
    r = cc.post("/api/auth/password", json={"old_password": "m_fresh123", "new_password": "m_fresh456"}, headers=th)
    assert r.json() == {"ok": True}
    me2 = cc.get("/api/auth/me", headers=th).json()
    assert me2["must_change_pw"] is False


def test_duplicate_resets_consent():
    cc, th, _ = _client()
    data = {"child_code": "D-1", "grade": "小4", "class_type": "通級",
            "profile_strengths": "a", "profile_needs": "b", "guardian_wish": "c",
            "support_long_goal": "d", "guidance_long_goal": "e", "short_goal_1": "f",
            "supports": "g", "eval_method": "h", "guardian_confirmed": True, "handover_consent": True}
    pid = cc.post("/api/plans", json={"child_code": "D-1", "grade": "小4", "class_type": "通級", "data": data}, headers=th).json()["id"]
    cc.post(f"/api/plans/{pid}/comments", json={"body": "申送り"}, headers=th)
    nid = cc.post(f"/api/plans/{pid}/duplicate", headers=th).json()["id"]
    got = cc.get(f"/api/plans/{nid}", headers=th).json()
    assert got["status"] == "draft"
    assert got["data"]["guardian_confirmed"] is False and got["data"]["handover_consent"] is False
    assert got["data"]["prev_plan_id"] == pid
    assert len(cc.get(f"/api/plans/{nid}/comments", headers=th).json()) == 1


def test_health_and_audit_export():
    cc, _, mh = _client()
    h = cc.get("/api/health").json()
    assert h["ok"] is True and "pdf_font" not in h
    d = cc.get("/api/health/detail", headers=mh).json()
    assert d["ok"] is True and d["pdf_font"] is True and d["db"] in ("sqlite", "postgres")
    r = cc.get("/api/audit/export", headers=mh)
    assert r.status_code == 200
    assert "操作" in r.content.decode("utf-8-sig")


def test_docx_append_and_table_leftover():
    from docx import Document
    from app.core import template_engine
    import tempfile
    tmp = Path(tempfile.mkdtemp())
    src = tmp / "t.docx"
    doc = Document()
    doc.add_paragraph("性別")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "{{missing_field}}"
    doc.save(src)
    out = tmp / "o.docx"
    # append モード
    template_engine.merge_docx(str(src), str(out),
                               {"gender": "女", "missing_field": "x"},
                               {"docx_p0": {"field": "gender", "mode": "append"}})
    assert Document(str(out)).paragraphs[0].text == "性別　女"
    # 表セル leftover 検査の確認：未解決があれば例外
    try:
        template_engine.merge_docx(str(src), str(tmp / "o2.docx"), {}, {})
        raise SystemExit("leftover未検出は異常")
    except ValueError as e:
        assert "missing_field" in str(e)


def test_backup_roundtrip():
    import shutil
    import tempfile
    import zipfile
    sys.path.insert(0, str(BASE / "backend"))
    import backup as backup_mod
    d = Path(tempfile.mkdtemp())
    z = backup_mod.backup(str(d))
    assert z.exists() and zipfile.is_zipfile(z)
    names = zipfile.ZipFile(z).namelist()
    assert any("app.db" in n for n in names)
    shutil.rmtree(d, ignore_errors=True)


def test_robust_merge_details(tmp_path):
    from openpyxl import Workbook
    from app.core import template_engine
    # アンダースコア付きシート名・数字セル保護
    p = tmp_path / 'u.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Sheet_A'
    ws['A1'] = '長期目標'; ws['B1'] = None; ws['C1'] = '4'
    wb.save(p)
    out = tmp_path / 'o.xlsx'
    template_engine.merge_xlsx(str(p), str(out), {'guidance_long_goal': '切替を行う'}, {'xlsx_Sheet_A_B1': 'guidance_long_goal', 'xlsx_Sheet_A_C1': 'guidance_long_goal'})
    from openpyxl import load_workbook
    r = load_workbook(out)['Sheet_A']
    assert r['B1'].value == '切替を行う'
    assert r['C1'].value == '4'
    # 日付定型セルは置換される
    p2 = tmp_path / 'd.xlsx'
    wb2 = Workbook(); ws2 = wb2.active; ws2['A1'] = '　　　　年　　月　　日'
    wb2.save(p2)
    out2 = tmp_path / 'd2.xlsx'
    template_engine.merge_xlsx(str(p2), str(out2), {'written_date': '令和8年4月'}, {'xlsx_Sheet_A1': 'written_date'})
    assert load_workbook(out2).active['A1'].value == '令和8年4月'


def test_logout_csvupload_guards():
    sys.path.insert(0, str(BASE / 'backend'))
    import io as _io
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    auth_mod.ensure_user('s_user', 's_user123', 'teacher', '')
    cc = TestClient(app)
    t = cc.post('/api/auth/login', json={'username': 's_user', 'password': 's_user123'}).json()['token']
    hh = {'Authorization': f'Bearer {t}'}
    # logoutでセッション無効化
    assert cc.post('/api/auth/logout', headers=hh).json() == {'ok': True}
    assert cc.get('/api/plans', headers=hh).status_code == 401
    # CSVインジェクション中和
    from app.routers.handover import csv_safe
    assert csv_safe('=cmd|1') == chr(39) + '=cmd|1'
    assert csv_safe('普通') == '普通'
    # 壊れたxlsxは400
    t2 = cc.post('/api/auth/login', json={'username': 's_user', 'password': 's_user123'}).json()['token']
    hh2 = {'Authorization': f'Bearer {t2}'}
    r = cc.post('/api/templates/upload', files={'file': ('x.xlsx', _io.BytesIO(b'not a zip'))}, headers=hh2)
    assert r.status_code == 400, r.status_code


def test_merge_prunes_old_outputs():
    sys.path.insert(0, str(BASE / 'backend'))
    import time as _t
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    from app.seed_data import register_niigata
    register_niigata()
    auth_mod.ensure_user('p_user', 'p_user123', 'teacher', '')
    cc = TestClient(app)
    t = cc.post('/api/auth/login', json={'username': 'p_user', 'password': 'p_user123'}).json()['token']
    hh = {'Authorization': f'Bearer {t}'}
    p = __import__('pathlib').Path('backend/data/templates/niigata01')
    data = {'child_code': 'P-1', 'grade': '小4', 'class_type': '通級', 'profile_strengths': 'a', 'profile_needs': 'b', 'guardian_wish': 'c', 'support_long_goal': 'd', 'guidance_long_goal': 'e', 'short_goal_1': 'f', 'supports': 'g', 'eval_method': 'h', 'guardian_confirmed': True}
    for _ in range(4):
        r = cc.post('/api/templates/niigata01/merge', json={'data': data}, headers=hh)
        assert r.status_code == 200, r.text[:200]
        _t.sleep(0.05)
    assert len(list(p.glob('filled_*.xlsx'))) <= 3


def test_viewer_blocked_and_input_caps():
    sys.path.insert(0, str(BASE / 'backend'))
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    from app.seed_data import register_niigata
    register_niigata()
    auth_mod.ensure_user('v_viewer', 'v_viewer123', 'viewer', '')
    auth_mod.ensure_user('v_teacher', 'v_teacher123', 'teacher', '')
    cc = TestClient(app)
    vt = cc.post('/api/auth/login', json={'username': 'v_viewer', 'password': 'v_viewer123'}).json()['token']
    vh = {'Authorization': f'Bearer {vt}'}
    tt = cc.post('/api/auth/login', json={'username': 'v_teacher', 'password': 'v_teacher123'}).json()['token']
    th = {'Authorization': f'Bearer {tt}'}
    pid = cc.post('/api/plans', json={'child_code': 'V-1', 'grade': '小4', 'class_type': '通級', 'data': {'profile_strengths': 'a'}}, headers=th).json()['id']
    # viewerは出力系不可
    assert cc.post('/api/templates/niigata01/merge', json={'plan_id': pid}, headers=vh).status_code == 403
    assert cc.get('/api/templates/niigata01/file/original.xlsx', headers=vh).status_code == 403
    assert cc.post('/api/templates/niigata01/pdf', json={'plan_id': pid}, headers=vh).status_code == 403
    assert cc.get(f'/api/plans/{pid}/handover?format=csv', headers=vh).status_code == 403
    # viewer閲覧は可
    assert cc.get(f'/api/plans/{pid}', headers=vh).status_code == 200
    # 入力上限
    big = {'child_code': 'V-2', 'grade': '小4', 'class_type': '通級', 'data': {'x': 'y' * (600 * 1024)}}
    assert cc.post('/api/plans', json=big, headers=th).status_code == 400
    assert cc.post('/api/auth/login', json={'username': 'x' * 300, 'password': 'p'}).status_code == 400


def test_ownership_boundary():
    sys.path.insert(0, str(BASE / 'backend'))
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    auth_mod.ensure_user('o_a', 'o_a1234567', 'teacher', '組A')
    auth_mod.ensure_user('o_b', 'o_b1234567', 'teacher', '組B')
    cc = TestClient(app)
    ta = cc.post('/api/auth/login', json={'username': 'o_a', 'password': 'o_a1234567'}).json()['token']
    tb = cc.post('/api/auth/login', json={'username': 'o_b', 'password': 'o_b1234567'}).json()['token']
    ha, hb = {'Authorization': f'Bearer {ta}'}, {'Authorization': f'Bearer {tb}'}
    full = {'child_code': 'O-1', 'grade': '小4', 'class_type': '通級', 'profile_strengths': 'a', 'profile_needs': 'b', 'guardian_wish': 'c', 'support_long_goal': 'd', 'guidance_long_goal': 'e', 'short_goal_1': 'f', 'supports': 'g', 'eval_method': 'h', 'guardian_confirmed': True}
    pid = cc.post('/api/plans', json={'child_code': 'O-1', 'grade': '小4', 'class_type': '通級', 'school': '組A', 'data': full}, headers=ha).json()['id']
    # Bからは参照・更新・状態・コメント・差し込み参照すべて403
    assert cc.get(f'/api/plans/{pid}', headers=hb).status_code == 403
    assert cc.put(f'/api/plans/{pid}', json={'child_code': 'O-1', 'data': full}, headers=hb).status_code == 403
    assert cc.post(f'/api/plans/{pid}/status', json={'status': 'review'}, headers=hb).status_code == 403
    assert cc.get(f'/api/plans/{pid}/comments', headers=hb).status_code == 403
    assert cc.post('/api/templates/niigata01/preview', json={'plan_id': pid}, headers=hb).status_code == 403
    assert cc.get(f'/api/plans/{pid}/handover?format=csv', headers=hb).status_code == 403
    # A本人は可
    assert cc.get(f'/api/plans/{pid}', headers=ha).status_code == 200


def test_caps_and_prune():
    sys.path.insert(0, str(BASE / 'backend'))
    from fastapi.testclient import TestClient
    from app.main import app
    from app import auth as auth_mod
    from app.seed_data import register_niigata
    register_niigata()
    auth_mod.ensure_user('cap_mgr', 'cap_mgr123', 'manager', '')
    cc = TestClient(app)
    t = cc.post('/api/auth/login', json={'username': 'cap_mgr', 'password': 'cap_mgr123'}).json()['token']
    mh = {'Authorization': f'Bearer {t}'}
    assert isinstance(cc.get('/api/audit?limit=999999', headers=mh).json(), list)
    r = cc.post('/api/templates/niigata01/mapping', json={'mapping': {f'k{i}': 'v' for i in range(2001)}}, headers=mh)
    assert r.status_code == 400, r.status_code
