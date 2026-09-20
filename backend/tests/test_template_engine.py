"""汎用エンジンの簡易テスト。実行: python -m pytest backend/tests -q"""
from pathlib import Path
import sys

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE / "backend"))
sys.path.insert(0, str(BASE / "src"))

from app.core import template_engine


def test_extract_txt_placeholder():
    slots = template_engine.extract_txt("目標 {{guidance_long_goal}} です")
    assert slots and slots[0]["field_hint"] == "guidance_long_goal"


def test_extract_xlsx_blank_with_label(tmp_path):
    from openpyxl import Workbook
    p = tmp_path / "t.xlsx"
    wb = Workbook(); ws = wb.active
    ws["A1"] = "長期目標"; ws["B1"] = None
    wb.save(p)
    slots = template_engine.extract_xlsx(str(p))
    assert any(s["slot_id"].endswith("B1") for s in slots)


def test_merge_xlsx(tmp_path):
    from openpyxl import Workbook
    p = tmp_path / "t.xlsx"
    wb = Workbook(); ws = wb.active
    ws["A1"] = "長期目標"; ws["B1"] = "{{guidance_long_goal}}"
    wb.save(p)
    out = tmp_path / "o.xlsx"
    template_engine.merge_xlsx(str(p), str(out), {"guidance_long_goal": "切替を行う"}, {})
    from openpyxl import load_workbook
    assert load_workbook(out).active["B1"].value == "切替を行う"
