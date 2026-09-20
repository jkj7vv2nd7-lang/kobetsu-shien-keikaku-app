"""汎用slot抽出＋差し込み v0.4。依存：python-docx, openpyxl（backend用）。"""
from __future__ import annotations
import re
from pathlib import Path

PLACEHOLDER = re.compile(r"\{\{(\w+)\}\}")
BLANK = re.compile(r"(＿{2,}|_{2,}|（\s*）|【\s*】|［\s*］|\(\s*\))")
# 日付等の定型空白セル（例：「　　年　　月　　日」「　　　年」「　(　　　)　　　-」）
FORM_BLANK = re.compile(r"^[\s　]*(\d*\s*[年年月日()（）\-―─－]*\s*)+$")


def extract_txt(text: str) -> list:
    slots = []
    for m in PLACEHOLDER.finditer(text):
        slots.append({"slot_id": f"txt_{m.group(1)}", "kind": "placeholder", "field_hint": m.group(1), "label": m.group(0)})
    return slots


def extract_docx(path: str) -> list:
    from docx import Document
    doc = Document(path)
    slots: list = []
    def add(slot_id, label, excerpt):
        slots.append({"slot_id": slot_id, "kind": "blank", "label": label[:40], "excerpt": excerpt[:80]})
    for i, p in enumerate(doc.paragraphs):
        for m in PLACEHOLDER.finditer(p.text):
            slots.append({"slot_id": f"docx_p{i}_{m.group(1)}", "kind": "placeholder",
                          "field_hint": m.group(1), "label": m.group(0)})
        if BLANK.search(p.text):
            add(f"docx_p{i}", p.text.strip() or f"段落{i}", p.text)
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                txt = cell.text or ""
                for m in PLACEHOLDER.finditer(txt):
                    slots.append({"slot_id": f"docx_t{ti}_r{ri}c{ci}_{m.group(1)}", "kind": "placeholder",
                                  "field_hint": m.group(1), "label": m.group(0)})
                if BLANK.search(txt) or (not txt.strip() and ri > 0):
                    hdr = ""
                    if ci > 0:
                        hdr = (row.cells[ci - 1].text or "").strip()
                    add(f"docx_t{ti}_r{ri}c{ci}", hdr or txt.strip() or f"表{ti}({ri},{ci})", txt)
    return slots


def extract_xlsx(path: str) -> list:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    slots: list = []
    for ws in wb.worksheets:
        max_r = max(ws.max_row or 1, 1)
        max_c = max(ws.max_column or 1, 1)
        for r in range(1, max_r + 2):  # +1で末尾の空欄候補も拾う
            for col in range(1, max_c + 2):
                c = ws.cell(row=r, column=col)
                v = c.value
                s = "" if v is None else str(v)
                for m in PLACEHOLDER.finditer(s):
                    slots.append({"slot_id": f"xlsx_{ws.title}_{c.coordinate}_{m.group(1)}",
                                  "kind": "placeholder", "field_hint": m.group(1),
                                  "label": m.group(0), "sheet": ws.title, "cell": c.coordinate})
                is_blankish = (v is None or (isinstance(v, str) and (not v.strip() or BLANK.search(v))))
                if is_blankish:
                    label = ""
                    if c.column > 1:
                        lv = ws.cell(row=c.row, column=c.column - 1).value
                        label = (str(lv).strip() if lv else "")
                    if not label and c.row > 1:
                        uv = ws.cell(row=c.row - 1, column=c.column).value
                        label = (str(uv).strip() if uv else "")
                    if label or BLANK.search(s):
                        slots.append({"slot_id": f"xlsx_{ws.title}_{c.coordinate}", "kind": "cell",
                                      "label": label or s[:40], "sheet": ws.title, "cell": c.coordinate})
    return slots


def merge_docx(src: str, dest: str, data: dict, mapping: dict) -> None:
    """mapping: slot_id -> field_id。placeholderは直接field名で置換。"""
    from docx import Document
    doc = Document(src)
    def _field(spec):
        return spec.get("field", "") if isinstance(spec, dict) else spec
    def _mode(spec):
        return spec.get("mode", "replace") if isinstance(spec, dict) else "replace"
    inv = {k: (data.get(_field(v), "")) for k, v in mapping.items()}
    modes = {k: _mode(v) for k, v in mapping.items()}
    def _apply(cur_text: str, sid: str) -> str:
        val = str(inv.get(sid) or "")
        if sid in inv and val and modes.get(sid) == "append" and cur_text.strip():
            return cur_text.rstrip() + "　" + val
        if sid in inv and val and BLANK.search(cur_text):
            return BLANK.sub(val, cur_text, count=1)
        if sid in inv and val and not cur_text.strip():
            return val
        return cur_text
    def fill(text: str, slot_base: str) -> str:
        # {{field}} 置換
        def _ph(m):
            return str(data.get(m.group(1), m.group(0)))
        text = PLACEHOLDER.sub(_ph, text)
        return text
    for i, p in enumerate(doc.paragraphs):
        sid = f"docx_p{i}"
        p.text = fill(_apply(p.text, sid), sid)
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    sid = f"docx_t{ti}_r{ri}c{ci}"
                    p.text = fill(_apply(p.text, sid), sid)
    # mapping漏れの{{}}残存検査（段落＋表セル）
    leftover = []
    for p in doc.paragraphs:
        leftover += PLACEHOLDER.findall(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                leftover += PLACEHOLDER.findall(cell.text)
    if leftover:
        raise ValueError(f"未解決プレースホルダ: {sorted(set(leftover))}")
    doc.save(dest)


def _is_writable_form_cell(cur) -> bool:
    """空・下線・{{}}・日付等の定型空白なら上書き可。"""
    if cur is None:
        return True
    if not isinstance(cur, str) or not cur.strip():
        return True
    if BLANK.search(cur) or PLACEHOLDER.search(cur):
        return True
    s = cur.strip()
    if s.isdigit():
        return False  # 数字のみの実データは保護
    return len(s) <= 20 and bool(FORM_BLANK.match(s))


def _fit_cell(ws, coord: str, text: str) -> None:
    """折り返し＋行高の簡易調整（長文のはみ出し対策）。"""
    from copy import copy
    try:
        c = ws[coord]
        al = copy(c.alignment)
        al.wrap_text = True
        al.vertical = "top"
        c.alignment = al
        n = len(str(text))
        if n > 30:
            row = ws.row_dimensions[c.row]
            need = 15 * (min(n, 300) // 30 + 1)
            if not row.height or row.height < need:
                row.height = min(need, 120)
    except Exception:
        pass


def merge_xlsx(src: str, dest: str, data: dict, mapping: dict) -> None:
    """mapping: slot_id -> field_id または {"field":..,"mode":"replace"|"append"}。
    appendは「既存ラベル＋値」で書き込む（ラベル一体型セル用）。"""
    from openpyxl import load_workbook
    wb = load_workbook(src)
    # 1) 明示マッピングを優先適用
    for slot_id, spec in mapping.items():
        # slot_id形式 xlsx_{sheet}_{cell} / xlsx_{sheet}_{cell}_{field}
        # シート名に"_"を含む場合に備え、既知シート名で前方一致させる
        if not slot_id.startswith("xlsx_"):
            continue
        rest_all = slot_id[len("xlsx_"):]
        sheet = next((sn for sn in wb.sheetnames if rest_all == sn or rest_all.startswith(sn + "_")), None)
        if sheet is None:
            continue
        rest = rest_all[len(sheet) + 1:] if rest_all != sheet else ""
        cell = rest.split("_")[0]  # B5 等（後ろにfield名が付く場合あり）
        if not cell:
            continue
        if isinstance(spec, dict):
            field_id, mode = spec.get("field", ""), spec.get("mode", "replace")
        else:
            field_id, mode = spec, "replace"
        ws = wb[sheet]
        val = data.get(field_id, "")
        if isinstance(val, bool):
            val = "確認済" if val else "未確認"
        val = "" if val is None else str(val)
        if not val:
            continue
        cur = ws[cell].value
        if mode == "append" and isinstance(cur, str) and cur.strip():
            ws[cell].value = cur.rstrip() + "　" + val
            _fit_cell(ws, cell, ws[cell].value)
        elif _is_writable_form_cell(cur):
            ws[cell].value = val
            _fit_cell(ws, cell, val)
        else:
            # placeholder埋め込み型
            new = PLACEHOLDER.sub(lambda m: str(data.get(m.group(1), m.group(0))), str(cur))
            if new != cur:
                ws[cell].value = new
                _fit_cell(ws, cell, new)
    # 2) マッピング未指定でも {{field}} はdataで直接置換（自治体様式の手間削減）
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and PLACEHOLDER.search(c.value):
                    c.value = PLACEHOLDER.sub(lambda m: str(data.get(m.group(1), m.group(0))), c.value)
    # {{}}残存検査
    left = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    left += PLACEHOLDER.findall(c.value)
    if left:
        raise ValueError(f"未解決プレースホルダ: {sorted(set(left))}")
    wb.save(dest)
